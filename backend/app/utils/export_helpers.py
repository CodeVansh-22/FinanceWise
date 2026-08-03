import io
import logging

logger = logging.getLogger(__name__)

def generate_transactions_excel(transactions):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = "Transactions"

        headers = ["ID", "Type", "Category", "Description", "Amount (INR)", "Date", "Status"]
        ws.append(headers)

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for t in transactions:
            ws.append([
                str(t.get("_id", t.get("id", ""))),
                t.get("type", "").upper(),
                t.get("category", ""),
                t.get("description", ""),
                float(t.get("amount", 0)),
                t.get("date", ""),
                t.get("status", "completed")
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    except Exception as e:
        logger.error(f"Excel generation error: {e}")
        # Fallback CSV format
        csv_str = "ID,Type,Category,Description,Amount,Date,Status\n"
        for t in transactions:
            csv_str += f"{t.get('_id','')},{t.get('type','')},{t.get('category','')},{t.get('description','')},{t.get('amount',0)},{t.get('date','')},{t.get('status','')}\n"
        return csv_str.encode('utf-8')

def generate_transactions_pdf(transactions, title="Financial Statement"):
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=20,
            textColor=colors.HexColor('#0F172A'),
            spaceAfter=12
        )
        elements.append(Paragraph(f"FinanceWise - {title}", title_style))
        elements.append(Spacer(1, 10))

        data = [["Date", "Type", "Category", "Description", "Amount (₹)"]]
        total_income = 0
        total_expense = 0

        for t in transactions:
            t_type = t.get("type", "expense")
            amt = float(t.get("amount", 0))
            if t_type == "income":
                total_income += amt
            else:
                total_expense += amt

            data.append([
                str(t.get("date", ""))[:10],
                t_type.upper(),
                str(t.get("category", "")),
                str(t.get("description", ""))[:30],
                f"₹{amt:,.2f}"
            ])

        data.append(["SUMMARY", "", "", f"Net: ₹{(total_income - total_expense):,.2f}", f"Inc: ₹{total_income:,.0f} | Exp: ₹{total_expense:,.0f}"])

        t_table = Table(data, colWidths=[80, 60, 100, 180, 110])
        t_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (4, 0), (4, -1), 'RIGHT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -2), colors.HexColor('#F8FAFC')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#F1F5F9')),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ]))
        elements.append(t_table)

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()
    except Exception as e:
        logger.error(f"PDF generation error: {e}")
        text_content = f"FinanceWise {title}\n"
        for t in transactions:
            text_content += f"{t.get('date')} | {t.get('type')} | {t.get('category')} | ₹{t.get('amount')}\n"
        return text_content.encode('utf-8')
